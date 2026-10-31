from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document as WordDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from docchat.chunking import split_documents
from docchat.loaders import load_document


@pytest.mark.parametrize(
    ("filename", "content", "expected"),
    [
        ("notes.txt", b"Plain text content", "Plain text content"),
        ("notes.md", b"# Heading\nMarkdown body", "Markdown body"),
    ],
)
def test_loads_text_formats(filename: str, content: bytes, expected: str) -> None:
    documents = load_document(filename, content)
    assert len(documents) == 1
    assert expected in documents[0].page_content
    assert documents[0].metadata["source"] == filename


def test_loads_docx_paragraphs_and_tables() -> None:
    buffer = BytesIO()
    word = WordDocument()
    word.add_paragraph("DOCX paragraph")
    word.add_table(rows=1, cols=1).cell(0, 0).text = "DOCX table"
    word.save(buffer)

    documents = load_document("report.docx", buffer.getvalue())

    assert "DOCX paragraph" in documents[0].page_content
    assert "DOCX table" in documents[0].page_content


def test_loads_pdf_with_page_metadata() -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): font_reference}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 220 Td (PDF page content) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)
    documents = load_document("report.pdf", buffer.getvalue())

    assert documents[0].page_content == "PDF page content"
    assert documents[0].metadata == {"source": "report.pdf", "page": 1}


def test_rejects_unsupported_and_invalid_text() -> None:
    with pytest.raises(ValueError, match="Unsupported"):
        load_document("image.png", b"bytes")
    with pytest.raises(ValueError, match="UTF-8"):
        load_document("bad.txt", b"\xff")


def test_chunking_preserves_metadata_and_overlap() -> None:
    from langchain_core.documents import Document

    source = Document(page_content="abcdefghij" * 4, metadata={"source": "a.pdf", "page": 2})
    chunks = split_documents([source], chunk_size=15, chunk_overlap=3)

    assert len(chunks) > 1
    assert all(chunk.metadata["source"] == "a.pdf" for chunk in chunks)
    assert all(chunk.metadata["page"] == 2 for chunk in chunks)
    assert chunks[0].page_content[-3:] == chunks[1].page_content[:3]
