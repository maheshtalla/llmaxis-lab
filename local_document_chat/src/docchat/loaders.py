"""Local document text extraction for supported upload formats."""

from __future__ import annotations

import io
from pathlib import Path

from langchain_core.documents import Document
from docx import Document as WordDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def load_document(filename: str, content: bytes) -> list[Document]:
    """Extract upload contents as LangChain documents with source metadata."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix or '(no extension)'}")
    if not content:
        raise ValueError(f"{filename} is empty.")

    if suffix == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(content))
            pages = [
                Document(
                    page_content=page.extract_text() or "",
                    metadata={"source": filename, "page": page_number},
                )
                for page_number, page in enumerate(reader.pages, start=1)
            ]
        except Exception as exc:
            raise ValueError(f"Could not read PDF {filename}: {exc}") from exc
        if not any(page.page_content.strip() for page in pages):
            raise ValueError(f"No extractable text was found in {filename}.")
        return pages

    if suffix == ".docx":
        try:
            word_doc = WordDocument(io.BytesIO(content))
            paragraphs = [paragraph.text for paragraph in word_doc.paragraphs]
            paragraphs.extend(
                cell.text
                for table in word_doc.tables
                for row in table.rows
                for cell in row.cells
            )
            text = "\n".join(part for part in paragraphs if part.strip())
        except Exception as exc:
            raise ValueError(f"Could not read DOCX {filename}: {exc}") from exc
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError(f"{filename} is not valid UTF-8 text.") from exc

    if not text.strip():
        raise ValueError(f"No extractable text was found in {filename}.")
    return [Document(page_content=text, metadata={"source": filename})]
